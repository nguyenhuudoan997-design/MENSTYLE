
"""Các route quản trị MENSTYLE."""

from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation

from sqlalchemy.exc import IntegrityError
from sqlalchemy import func


from flask import (
    Blueprint,
    abort,
    flash,
    redirect,
    render_template,
    request,
    url_for,
)
from flask_login import current_user, login_required

from app.extensions import db
from app.models import Order, Payment, Product, ProductSize, User
from werkzeug.security import generate_password_hash
from app.services.audit_service import log_action


from io import BytesIO

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


admin_bp = Blueprint("admin", __name__, url_prefix="/admin")


def admin_required():
    """Chỉ cho phép tài khoản admin truy cập."""
    if current_user.role != "admin":
        abort(403)



@admin_bp.route("/")
@login_required
def dashboard():
    """Trang tổng quan quản trị với thống kê doanh thu và đơn hàng."""
    admin_required()

    # Đọc bộ lọc ngày từ URL.
    start_date_text = request.args.get("start_date", "").strip()
    end_date_text = request.args.get("end_date", "").strip()

    start_date = None
    end_date = None
    date_error = None

    try:
        if start_date_text:
            start_date = datetime.strptime(
                start_date_text, "%Y-%m-%d"
            ).date()

        if end_date_text:
            end_date = datetime.strptime(
                end_date_text, "%Y-%m-%d"
            ).date()

        if start_date and end_date and start_date > end_date:
            date_error = "Ngày bắt đầu không được sau ngày kết thúc."

    except ValueError:
        date_error = "Ngày không hợp lệ. Vui lòng chọn lại."

    # Không áp dụng bộ lọc nếu dữ liệu ngày không hợp lệ.
    if date_error:
        start_date = None
        end_date = None

    # Thống kê đơn hàng theo trạng thái.
    total_orders = Order.query.count()

    pending_orders = Order.query.filter_by(
        order_status="Chờ xác nhận"
    ).count()

    shipping_orders = Order.query.filter_by(
        order_status="Đang giao"
    ).count()

    completed_orders = Order.query.filter_by(
        order_status="Hoàn thành"
    ).count()

    cancelled_orders = Order.query.filter_by(
        order_status="Đã hủy"
    ).count()

    # Doanh thu: chỉ tính thanh toán thành công của đơn không bị hủy.
    revenue_query = (
        db.session.query(func.coalesce(func.sum(Payment.amount), 0))
        .join(Order, Payment.order_id == Order.id)
        .filter(
            Payment.payment_status == "paid",
            Order.order_status != "Đã hủy",
        )
    )

    # Lọc theo thời điểm thanh toán thực tế.
    if start_date:
        revenue_query = revenue_query.filter(
            Payment.paid_at >= datetime.combine(
                start_date, datetime.min.time()
            )
        )

    if end_date:
        revenue_query = revenue_query.filter(
            Payment.paid_at < datetime.combine(
                end_date + timedelta(days=1),
                datetime.min.time(),
            )
        )

    total_revenue = revenue_query.scalar() or Decimal("0")

    return render_template(
        "admin/dashboard.html",
        total_orders=total_orders,
        pending_orders=pending_orders,
        shipping_orders=shipping_orders,
        completed_orders=completed_orders,
        cancelled_orders=cancelled_orders,
        total_revenue=total_revenue,
        start_date=start_date_text,
        end_date=end_date_text,
        date_error=date_error,
    )

@admin_bp.route("/reports/revenue")
@login_required
def revenue_report():
    """Báo cáo doanh thu theo khoảng thời gian."""
    admin_required()

    # Đọc bộ lọc thời gian.
    from datetime import date

    period = request.args.get("period", "").strip()
    start_text = request.args.get("start_date", "").strip()
    end_text = request.args.get("end_date", "").strip()

    today = date.today()
    start_date = None
    end_date = None
    date_error = None

    # Các bộ lọc nhanh.
    if period == "today":
        start_date = today
        end_date = today

    elif period == "last7":
        start_date = today - timedelta(days=6)
        end_date = today

    elif period == "this_month":
        start_date = today.replace(day=1)
        end_date = today

    elif period == "this_year":
        start_date = today.replace(month=1, day=1)
        end_date = today

    elif period:
        date_error = "Bộ lọc thời gian không hợp lệ."
        period = ""

    # Nếu không chọn bộ lọc nhanh, đọc ngày người dùng nhập.
    if not period and not date_error:
        try:
            if start_text:
                start_date = datetime.strptime(
                    start_text, "%Y-%m-%d"
                ).date()

            if end_text:
                end_date = datetime.strptime(
                    end_text, "%Y-%m-%d"
                ).date()

            if start_date and end_date and start_date > end_date:
                date_error = "Ngày bắt đầu không được sau ngày kết thúc."

        except ValueError:
            date_error = "Ngày không hợp lệ. Vui lòng chọn lại."

    # Chỉ giữ bộ lọc hợp lệ.
    if date_error:
        start_date = None
        end_date = None
        start_text = ""
        end_text = ""
        period = ""
    else:
        start_text = start_date.isoformat() if start_date else ""
        end_text = end_date.isoformat() if end_date else ""


    start_date = None
    end_date = None
    date_error = None

    try:
        if start_text:
            start_date = datetime.strptime(
                start_text, "%Y-%m-%d"
            ).date()

        if end_text:
            end_date = datetime.strptime(
                end_text, "%Y-%m-%d"
            ).date()

        if start_date and end_date and start_date > end_date:
            date_error = "Ngày bắt đầu không được sau ngày kết thúc."

    except ValueError:
        date_error = "Ngày không hợp lệ. Vui lòng chọn lại."

    if date_error:
        start_date = None
        end_date = None

    # Truy vấn thanh toán thành công, không tính đơn đã hủy.
    revenue_query = (
        db.session.query(Payment)
        .join(Order, Payment.order_id == Order.id)
        .filter(
            Payment.payment_status == "paid",
            Payment.paid_at.isnot(None),
            Order.order_status != "Đã hủy",
        )
    )

    if start_date:
        revenue_query = revenue_query.filter(
            Payment.paid_at >= datetime.combine(
                start_date, datetime.min.time()
            )
        )

    if end_date:
        revenue_query = revenue_query.filter(
            Payment.paid_at < datetime.combine(
                end_date + timedelta(days=1),
                datetime.min.time(),
            )
        )

    # Lấy các khoản thanh toán phù hợp với bộ lọc.
    payments = revenue_query.order_by(
        Payment.paid_at.desc(),
        Payment.id.desc(),
    ).all()
    
    # Gom doanh thu theo ngày và theo tháng.
    daily_totals = {}
    monthly_totals = {}

    for payment in payments:
        paid_date = payment.paid_at.date()
        day_key = paid_date.strftime("%Y-%m-%d")
        month_key = paid_date.strftime("%Y-%m")

        amount = float(payment.amount)

        daily_totals[day_key] = (
            daily_totals.get(day_key, 0) + amount
        )
        monthly_totals[month_key] = (
            monthly_totals.get(month_key, 0) + amount
        )

    # Dữ liệu cho biểu đồ doanh thu theo ngày.
    daily_labels = sorted(daily_totals.keys())
    daily_values = [
        round(daily_totals[day], 2)
        for day in daily_labels
    ]

    # Dữ liệu cho biểu đồ doanh thu theo tháng.
    monthly_labels = sorted(monthly_totals.keys())
    monthly_values = [
        round(monthly_totals[month], 2)
        for month in monthly_labels
    ]


    total_revenue = sum(
        (payment.amount for payment in payments),
        Decimal("0"),
    )

    paid_order_ids = {payment.order_id for payment in payments}
    paid_order_count = len(paid_order_ids)

    # Đếm đơn hoàn thành trong khoảng thời gian đã chọn.
    completed_query = Order.query.filter_by(
        order_status="Hoàn thành"
    )

    if start_date:
        completed_query = completed_query.filter(
            Order.created_at >= datetime.combine(
                start_date, datetime.min.time()
            )
        )

    if end_date:
        completed_query = completed_query.filter(
            Order.created_at < datetime.combine(
                end_date + timedelta(days=1),
                datetime.min.time(),
            )
        )

    completed_orders = completed_query.count()

    average_order_value = (
        total_revenue / paid_order_count
        if paid_order_count
        else Decimal("0")
    )

    return render_template(
        "admin/revenue_report.html",
        payments=payments,
        total_revenue=total_revenue,
        paid_order_count=paid_order_count,
        completed_orders=completed_orders,
        average_order_value=average_order_value,
        start_date=start_text,
        end_date=end_text,
        date_error=date_error,
        daily_labels=daily_labels,
        daily_values=daily_values,
        monthly_labels=monthly_labels,
        monthly_values=monthly_values,
        period=period,
    )


@admin_bp.route("/products")
@login_required
def product_list():
    """Hiển thị danh sách sản phẩm."""
    admin_required()

    products = (
        Product.query
        .filter_by(is_active=True)
        .order_by(Product.id.desc())
        .all()
    )

    return render_template(
        "admin/products.html",
        products=products,
    )


@admin_bp.route("/products/add", methods=["GET", "POST"])
@login_required
def product_add():
    """Thêm sản phẩm mới."""
    admin_required()

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        category = request.form.get("category", "").strip()
        material = request.form.get("material", "").strip()
        image = request.form.get("image", "").strip()
        description = request.form.get("description", "").strip()
        sizes = request.form.get("sizes", "").strip()
        colors = request.form.get("colors", "").strip()

        try:
            price = Decimal(request.form.get("price", "").strip())
            sale_price_text = request.form.get(
                "sale_price", ""
            ).strip()
            sale_price = (
                Decimal(sale_price_text)
                if sale_price_text
                else None
            )
            stock = int(request.form.get("stock", "").strip())
        except (InvalidOperation, ValueError):
            flash("Giá và số lượng tồn kho phải là số hợp lệ.", "error")
            return render_template("admin/product_form.html")

        if not name or not category:
            flash("Vui lòng nhập tên sản phẩm và danh mục.", "error")
            return render_template("admin/product_form.html")

        if price < 0 or (sale_price is not None and sale_price < 0):
            flash("Giá sản phẩm không được âm.", "error")
            return render_template("admin/product_form.html")

        if sale_price is not None and sale_price > price:
            flash("Giá khuyến mãi không được lớn hơn giá gốc.", "error")
            return render_template("admin/product_form.html")

        if stock < 0:
            flash("Số lượng tồn kho không được âm.", "error")
            return render_template("admin/product_form.html")

        product = Product(
            name=name,
            category=category,
            price=price,
            sale_price=sale_price,
            image=image or None,
            description=description or None,
            material=material or None,
            sizes=sizes or None,
            colors=colors or None,
            stock=stock,
            created_at=datetime.now(),
        )
      
        try:
            db.session.add(product)

            # Lưu sản phẩm trước để lấy ID do MySQL tạo.
            db.session.flush()

            # Ghi nhật ký trong cùng giao dịch với sản phẩm.
            log_action(
                action="CREATE",
                target_type="Product",
                target_id=product.id,
                description=f"Thêm sản phẩm: {product.name}",
            )

            # Lưu cả sản phẩm và nhật ký cùng lúc.
            db.session.commit()

        except Exception:
            db.session.rollback()
            admin_bp.logger.exception("Không thể lưu sản phẩm mới.")
            flash(
                "Không thể lưu sản phẩm. Hãy kiểm tra dữ liệu và kết nối MySQL.",
                "error",
            )
            return render_template("admin/product_form.html")


        flash("Đã thêm sản phẩm thành công.", "success")
        return redirect(url_for("admin.product_list"))

    return render_template("admin/product_form.html")

@admin_bp.route("/products/edit/<int:product_id>", methods=["GET", "POST"])
@login_required
def product_edit(product_id):
    """Chỉnh sửa thông tin sản phẩm."""
    admin_required()

    product = db.session.get(Product, product_id)

    if product is None:
        abort(404)

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        category = request.form.get("category", "").strip()
        material = request.form.get("material", "").strip()
        image = request.form.get("image", "").strip()
        description = request.form.get("description", "").strip()
        sizes = request.form.get("sizes", "").strip()
        colors = request.form.get("colors", "").strip()

        try:
            price = Decimal(request.form.get("price", "").strip())
            sale_price_text = request.form.get("sale_price", "").strip()
            sale_price = (
                Decimal(sale_price_text) if sale_price_text else None
            )
            stock = int(request.form.get("stock", "").strip())
        except (InvalidOperation, ValueError):
            flash("Giá và số lượng tồn kho phải là số hợp lệ.", "error")
            return render_template(
                "admin/product_form.html", product=product
            )

        if not name or not category:
            flash("Vui lòng nhập tên sản phẩm và danh mục.", "error")
        elif price < 0 or (sale_price is not None and sale_price < 0):
            flash("Giá sản phẩm không được âm.", "error")
        elif sale_price is not None and sale_price > price:
            flash("Giá khuyến mãi không được lớn hơn giá gốc.", "error")
        elif stock < 0:
            flash("Số lượng tồn kho không được âm.", "error")
        else:
            product.name = name
            product.category = category
            product.price = price
            product.sale_price = sale_price
            product.material = material or None
            product.image = image or None
            product.description = description or None
            product.sizes = sizes or None
            product.colors = colors or None
            product.stock = stock
            
            try:
                log_action(
                    action="UPDATE",
                    target_type="Product",
                    target_id=product.id,
                    description=f"Cập nhật sản phẩm: {product.name}",
                )

                db.session.commit()

                flash("Đã cập nhật sản phẩm thành công.", "success")
                return redirect(url_for("admin.product_list"))

            except Exception:
                db.session.rollback()
                admin_bp.logger.exception("Không thể cập nhật sản phẩm.")
                flash(
                    "Không thể cập nhật sản phẩm. Vui lòng thử lại.",
                    "error",
                )

    return render_template("admin/product_form.html", product=product)

@admin_bp.route("/products/delete/<int:product_id>", methods=["POST"])
@login_required
def product_delete(product_id):
    """Ngừng kinh doanh sản phẩm để bảo toàn lịch sử đơn hàng."""
    admin_required()

    product = db.session.get(Product, product_id)

    if product is None:
        abort(404)

    try:
        product.is_active = False

        log_action(
            action="UPDATE",
            target_type="Product",
            target_id=product.id,
            description=(
                f"Ngừng kinh doanh sản phẩm: {product.name}"
            ),
        )

        db.session.commit()

        flash(
            "Đã ngừng kinh doanh sản phẩm. "
            "Lịch sử đơn hàng được giữ nguyên.",
            "success",
        )

    except Exception:
        db.session.rollback()
        admin_bp.logger.exception(
            "Không thể ngừng kinh doanh sản phẩm."
        )
        flash(
            "Có lỗi xảy ra khi ngừng kinh doanh sản phẩm.",
            "error",
        )

    return redirect(url_for("admin.product_list"))


# =========================
# QUAN LY BANG SIZE
# =========================

SIZE_FIELDS = (
    "chest_cm",
    "waist_cm",
    "hip_cm",
    "shoulder_cm",
    "length_cm",
    "height_min_cm",
    "height_max_cm",
    "weight_min_kg",
    "weight_max_kg",
)

def parse_size_form():
    """Đọc và kiểm tra dữ liệu từ biểu mẫu size."""
    size_name = request.form.get("size_name", "").strip()

    if not size_name or len(size_name) > 20:
        raise ValueError("Tên size phải có từ 1 đến 20 ký tự.")

    values = {}

    for field in SIZE_FIELDS:
        raw_value = request.form.get(field, "").strip()

        if not raw_value:
            values[field] = None
            continue

        try:
            value = Decimal(raw_value)
        except InvalidOperation:
            raise ValueError(f"Giá trị {field} không hợp lệ.")

        if not value.is_finite() or value <= 0 or value > 1000:
            raise ValueError(f"Giá trị {field} phải lớn hơn 0 và không quá 1000.")

        values[field] = value

    # Kiểm tra khoảng chiều cao và cân nặng.
    if (
        values["height_min_cm"] is not None
        and values["height_max_cm"] is not None
        and values["height_min_cm"] > values["height_max_cm"]
    ):
        raise ValueError("Chiều cao tối thiểu không được lớn hơn tối đa.")

    if (
        values["weight_min_kg"] is not None
        and values["weight_max_kg"] is not None
        and values["weight_min_kg"] > values["weight_max_kg"]
    ):
        raise ValueError("Cân nặng tối thiểu không được lớn hơn tối đa.")

    return size_name, values


@admin_bp.route("/product-sizes")
@login_required
def product_size_list():
    """Hiển thị bảng size theo sản phẩm."""
    admin_required()

    products = Product.query.order_by(Product.id.desc()).all()
    product_id = request.args.get("product_id", type=int)

    query = ProductSize.query

    if product_id is not None:
        query = query.filter_by(product_id=product_id)

    size_charts = query.order_by(
        ProductSize.product_id.asc(),
        ProductSize.id.asc(),
    ).all()

    return render_template(
        "admin/product_sizes.html",
        products=products,
        size_charts=size_charts,
        selected_product_id=product_id,
    )


@admin_bp.route("/product-sizes/add", methods=["GET", "POST"])
@login_required
def product_size_add():
    """Thêm một size cho sản phẩm."""
    admin_required()

    products = Product.query.order_by(Product.name.asc()).all()
    selected_product_id = request.args.get("product_id", type=int)

    if request.method == "POST":
        selected_product_id = request.form.get(
            "product_id", type=int
        )

        product = db.session.get(Product, selected_product_id)

        if product is None:
            flash("Sản phẩm không tồn tại.", "error")
            return render_template(
                "admin/product_size_form.html",
                products=products,
                size=None,
                selected_product_id=selected_product_id,
            )

        try:
            size_name, values = parse_size_form()

            existing_size = ProductSize.query.filter_by(
                product_id=product.id,
                size_name=size_name,
            ).first()

            if existing_size:
                raise ValueError(
                    "Size này đã tồn tại trong sản phẩm được chọn."
                )

            size = ProductSize(
                product_id=product.id,
                size_name=size_name,
                **values,
            )

            db.session.add(size)
            db.session.flush()

            log_action(
                action="CREATE",
                target_type="ProductSize",
                target_id=size.id,
                description=(
                    f"Thêm size {size_name} cho sản phẩm {product.name}"
                ),
            )

            db.session.commit()

            flash("Đã thêm size thành công.", "success")
            return redirect(
                url_for(
                    "admin.product_size_list",
                    product_id=product.id,
                )
            )

        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "error")

        except IntegrityError:
            db.session.rollback()
            flash(
                "Size đã tồn tại cho sản phẩm này.",
                "error",
            )

        except Exception:
            db.session.rollback()
            admin_bp.logger.exception("Lỗi thêm bảng size.")
            flash("Không thể thêm size. Vui lòng thử lại.", "error")

    return render_template(
        "admin/product_size_form.html",
        products=products,
        size=None,
        selected_product_id=selected_product_id,
    )


@admin_bp.route(
    "/product-sizes/<int:size_id>/edit",
    methods=["GET", "POST"],
)
@login_required
def product_size_edit(size_id):
    """Chỉnh sửa thông số của một size."""
    admin_required()

    size = db.session.get(ProductSize, size_id)

    if size is None:
        abort(404)

    products = Product.query.order_by(Product.name.asc()).all()

    if request.method == "POST":
        product_id = request.form.get("product_id", type=int)
        product = db.session.get(Product, product_id)

        try:
            if product is None:
                raise ValueError("Vui lòng chọn sản phẩm hợp lệ.")

            size_name, values = parse_size_form()

            duplicate = ProductSize.query.filter(
                ProductSize.product_id == product.id,
                ProductSize.size_name == size_name,
                ProductSize.id != size.id,
            ).first()

            if duplicate:
                raise ValueError(
                    "Size này đã tồn tại trong sản phẩm được chọn."
                )

            old_product_name = size.product.name

            size.product_id = product.id
            size.size_name = size_name

            for field, value in values.items():
                setattr(size, field, value)

            log_action(
                action="UPDATE",
                target_type="ProductSize",
                target_id=size.id,
                description=(
                    f"Cập nhật size {size_name}; "
                    f"sản phẩm trước: {old_product_name}; "
                    f"sản phẩm hiện tại: {product.name}"
                ),
            )

            db.session.commit()

            flash("Đã cập nhật size thành công.", "success")
            return redirect(
                url_for(
                    "admin.product_size_list",
                    product_id=product.id,
                )
            )

        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "error")

        except IntegrityError:
            db.session.rollback()
            flash("Size bị trùng với dữ liệu đã có.", "error")

        except Exception:
            db.session.rollback()
            admin_bp.logger.exception("Lỗi cập nhật bảng size.")
            flash("Không thể cập nhật size. Vui lòng thử lại.", "error")

    return render_template(
        "admin/product_size_form.html",
        products=products,
        size=size,
        selected_product_id=size.product_id,
    )


@admin_bp.route(
    "/product-sizes/<int:size_id>/delete",
    methods=["POST"],
)
@login_required
def product_size_delete(size_id):
    """Xóa một size sau khi xác nhận."""
    admin_required()

    size = db.session.get(ProductSize, size_id)

    if size is None:
        abort(404)

    product_id = size.product_id
    size_name = size.size_name
    product_name = size.product.name

    try:
        log_action(
            action="DELETE",
            target_type="ProductSize",
            target_id=size.id,
            description=(
                f"Xóa size {size_name} của sản phẩm {product_name}"
            ),
        )

        db.session.delete(size)
        db.session.commit()

        flash("Đã xóa size thành công.", "success")

    except Exception:
        db.session.rollback()
        admin_bp.logger.exception("Lỗi xóa bảng size.")
        flash("Không thể xóa size. Vui lòng thử lại.", "error")

    return redirect(
        url_for(
            "admin.product_size_list",
            product_id=product_id,
        )
    )


# Những trạng thái được phép chuyển tiếp.
ORDER_TRANSITIONS = {
    "Chờ xác nhận": {"Chờ lấy hàng", "Đã hủy"},
    "Chờ lấy hàng": {"Đang giao", "Đã hủy"},
    "Đang giao": {"Đã giao"},
    "Đã giao": {"Hoàn thành"},
    "Hoàn thành": set(),
    "Đã hủy": set(),
}

ORDER_STATUSES = list(ORDER_TRANSITIONS.keys())



@admin_bp.route("/orders")
@login_required
def order_list():
    """Tìm kiếm, lọc và phân trang danh sách đơn hàng."""
    admin_required()

    # Lấy điều kiện tìm kiếm từ URL.
    keyword = request.args.get("keyword", "").strip()
    order_status = request.args.get("order_status", "").strip()
    payment_method = request.args.get("payment_method", "").strip()
    payment_status = request.args.get("payment_status", "").strip()

    # Các giá trị lọc được phép.
    allowed_order_statuses = ORDER_STATUSES
    allowed_payment_methods = {"COD", "Momo", "VNPay"}
    allowed_payment_statuses = {"pending", "paid", "failed"}

    # Khởi tạo truy vấn đơn hàng.
    query = Order.query

    # Tìm theo mã đơn, tên người nhận hoặc số điện thoại.
    if keyword:
        search_term = f"%{keyword}%"

        query = query.filter(
            db.or_(
                Order.order_code.ilike(search_term),
                Order.recipient_name.ilike(search_term),
                Order.recipient_phone.ilike(search_term),
            )
        )

    # Lọc trạng thái đơn hàng.
    if order_status:
        if order_status in allowed_order_statuses:
            query = query.filter(
                Order.order_status == order_status
            )
        else:
            flash("Trạng thái đơn hàng không hợp lệ.", "warning")
            order_status = ""

    # Lọc phương thức thanh toán.
    if payment_method:
        if payment_method in allowed_payment_methods:
            query = query.filter(
                Order.payment_method == payment_method
            )
        else:
            flash("Phương thức thanh toán không hợp lệ.", "warning")
            payment_method = ""

    # Lọc trạng thái thanh toán.
    if payment_status:
        if payment_status in allowed_payment_statuses:
            query = query.filter(
                Order.payment_status == payment_status
            )
        else:
            flash("Trạng thái thanh toán không hợp lệ.", "warning")
            payment_status = ""

    # Phân trang: mỗi trang tối đa 10 đơn hàng.
    page = request.args.get("page", 1, type=int)
    page = max(page or 1, 1)

    pagination = (
        query
        .order_by(Order.created_at.desc(), Order.id.desc())
        .paginate(
            page=page,
            per_page=10,
            error_out=False,
        )
    )

    return render_template(
        "admin/orders.html",
        orders=pagination.items,
        pagination=pagination,
        keyword=keyword,
        order_status=order_status,
        payment_method=payment_method,
        payment_status=payment_status,
        order_statuses=ORDER_STATUSES,
    )


@admin_bp.route("/orders/<int:order_id>")
@login_required
def order_detail(order_id):
    """Hiển thị chi tiết một đơn hàng."""
    admin_required()

    order = db.session.get(Order, order_id)

    if order is None:
        abort(404)

    return render_template(
        "admin/order_detail.html",
        order=order,
        order_statuses=ORDER_STATUSES,
        order_transitions=ORDER_TRANSITIONS,
    )



@admin_bp.route(
    "/orders/<int:order_id>/status",
    methods=["POST"],
)
@login_required
def order_update_status(order_id):
    """Cập nhật trạng thái và hoàn tồn kho khi hủy đơn hợp lệ."""
    admin_required()

    from sqlalchemy import select

    try:
        # Khóa đơn hàng trong giao dịch để tránh xử lý đồng thời.
        order = db.session.scalar(
            select(Order)
            .where(Order.id == order_id)
            .with_for_update()
        )

        if order is None:
            abort(404)

        new_status = request.form.get("order_status", "").strip()
        current_status = order.order_status

        if new_status not in ORDER_STATUSES:
            db.session.rollback()
            flash("Trạng thái đơn hàng không hợp lệ.", "error")
            return redirect(
                url_for("admin.order_detail", order_id=order_id)
            )

        # Không làm gì nếu trạng thái không thay đổi.
        if new_status == current_status:
            db.session.rollback()
            flash("Đơn hàng đang ở trạng thái này.", "warning")
            return redirect(
                url_for("admin.order_detail", order_id=order_id)
            )

        # Chỉ chấp nhận chuyển trạng thái theo quy trình.
        allowed_statuses = ORDER_TRANSITIONS.get(current_status, set())

        if new_status not in allowed_statuses:
            db.session.rollback()
            flash(
                f"Không thể chuyển từ '{current_status}' "
                f"sang '{new_status}'.",
                "error",
            )
            return redirect(
                url_for("admin.order_detail", order_id=order_id)
            )

        # Khi hủy đơn, hoàn lại số lượng từng sản phẩm đúng một lần.
        if new_status == "Đã hủy":
            items = list(order.items)
            product_ids = sorted({
                item.product_id for item in items
            })

            products = db.session.scalars(
                select(Product)
                .where(Product.id.in_(product_ids))
                .order_by(Product.id)
                .with_for_update()
            ).all()

            products_by_id = {
                product.id: product for product in products
            }

            for item in items:
                product = products_by_id.get(item.product_id)

                if product is None:
                    raise ValueError(
                        f"Không tìm thấy sản phẩm ID {item.product_id}; "
                        "không thể hoàn tồn kho."
                    )

            # Chỉ cộng kho sau khi đã xác nhận đủ sản phẩm.
            for item in items:
                products_by_id[item.product_id].stock += item.quantity

        # Lưu trạng thái mới của đơn hàng.
        order.order_status = new_status

        # Ghi nhật ký thay đổi trạng thái.
        log_action(
            action="UPDATE",
            target_type="Order",
            target_id=order.id,
            description=(
                f"Cập nhật trạng thái đơn {order.order_code}: "
                f"{current_status} -> {new_status}"
            ),
        )
        
        # Lưu trạng thái và nhật ký trong cùng giao dịch.
        db.session.commit()

        flash("Đã cập nhật trạng thái đơn hàng.", "success")

    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), "error")

    except Exception:
        db.session.rollback()
        admin_bp.logger.exception(
            "Không thể cập nhật trạng thái đơn hàng."
        )
        flash("Có lỗi xảy ra. Vui lòng thử lại.", "error")

    return redirect(
        url_for("admin.order_detail", order_id=order_id)
    )

@admin_bp.route(
    "/orders/<int:order_id>/confirm-cod-payment",
    methods=["POST"],
)
@login_required
def confirm_cod_payment(order_id):
    """Xác nhận đã thu tiền COD và cập nhật bản ghi thanh toán."""
    admin_required()

    from sqlalchemy import select
    from datetime import datetime
    from app.models import Payment

    try:
        order = db.session.scalar(
            select(Order)
            .where(Order.id == order_id)
            .with_for_update()
        )

        if order is None:
            abort(404)

        if order.payment_method != "COD":
            db.session.rollback()
            flash(
                "Chỉ xác nhận thu tiền bằng chức năng COD.",
                "error",
            )
            return redirect(
                url_for("admin.order_detail", order_id=order_id)
            )

        if order.order_status not in {"Đã giao", "Hoàn thành"}:
            db.session.rollback()
            flash(
                "Chỉ xác nhận thu tiền COD sau khi đơn đã giao.",
                "warning",
            )
            return redirect(
                url_for("admin.order_detail", order_id=order_id)
            )

        if order.payment_status == "paid":
            db.session.rollback()
            flash("Đơn hàng đã được xác nhận thanh toán.", "warning")
            return redirect(
                url_for("admin.order_detail", order_id=order_id)
            )

        payment = db.session.scalar(
            select(Payment)
            .where(Payment.order_id == order.id)
            .order_by(Payment.id.desc())
            .with_for_update()
        )

        if payment is None:
            raise ValueError(
                "Không tìm thấy bản ghi thanh toán của đơn hàng."
            )

        if payment.payment_status == "paid":
            raise ValueError(
                "Bản ghi thanh toán đã được xác nhận trước đó."
            )

        paid_time = datetime.now()

        order.payment_status = "paid"
        payment.payment_status = "paid"
        payment.paid_at = paid_time

        db.session.commit()
        flash("Đã xác nhận thu tiền COD thành công.", "success")

    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), "error")

    except Exception:
        db.session.rollback()
        admin_bp.logger.exception(
            "Không thể xác nhận thanh toán COD."
        )
        flash("Có lỗi khi xác nhận thanh toán.", "error")

    return redirect(
        url_for("admin.order_detail", order_id=order_id)
    )

# =========================
# QUAN LY NHAN VIEN
# =========================

@admin_bp.route("/employees")
@login_required
def employee_list():
    """Hien thi danh sach tai khoan nhan vien."""
    admin_required()

    employees = (
        User.query
        .filter_by(role="employee")
        .order_by(User.id.desc())
        .all()
    )

    return render_template(
        "admin/employees.html",
        employees=employees,
    )


@admin_bp.route("/employees/add", methods=["GET", "POST"])
@login_required
def employee_add():
    """Admin tao tai khoan nhan vien."""
    admin_required()

    if request.method == "POST":
        full_name = request.form.get("full_name", "").strip()
        email = request.form.get("email", "").strip().lower()
        phone = request.form.get("phone", "").strip()
        password = request.form.get("password", "")

        if not full_name or len(full_name) > 150:
            flash("Vui long nhap ho ten hop le.", "error")
            return render_template("admin/employee_form.html")

        if (
            not email
            or len(email) > 150
            or "@" not in email
            or any(char.isspace() for char in email)
        ):
            flash("Email khong hop le.", "error")
            return render_template("admin/employee_form.html")

        if phone and len(phone) > 20:
            flash("So dien thoai khong duoc qua 20 ky tu.", "error")
            return render_template("admin/employee_form.html")

        if len(password) < 8:
            flash("Mat khau phai co it nhat 8 ky tu.", "error")
            return render_template("admin/employee_form.html")

        existing_user = db.session.scalar(
            db.select(User).where(
                db.func.lower(User.email) == email
            )
        )

        if existing_user:
            flash("Email da ton tai.", "error")
            return render_template("admin/employee_form.html")

        employee = User(
            full_name=full_name,
            email=email,
            phone=phone or None,
            password_hash=generate_password_hash(password),
            role="employee",
            is_active=True,
        )

        try:
            db.session.add(employee)
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash("Email da ton tai hoac du lieu khong hop le.", "error")
            return render_template("admin/employee_form.html")
        except Exception:
            db.session.rollback()
            admin_bp.logger.exception("Loi tao tai khoan nhan vien.")
            flash("Khong the tao tai khoan nhan vien.", "error")
            return render_template("admin/employee_form.html")

        flash("Da tao tai khoan nhan vien.", "success")
        return redirect(url_for("admin.employee_list"))

    return render_template("admin/employee_form.html")


@admin_bp.route(
    "/employees/<int:employee_id>/toggle-active",
    methods=["POST"],
)
@login_required
def employee_toggle_active(employee_id):
    """Khoa hoac mo khoa tai khoan nhan vien."""
    admin_required()

    employee = db.session.get(User, employee_id)

    if employee is None or employee.role != "employee":
        abort(404)

    employee.is_active = not employee.is_active

    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        admin_bp.logger.exception(
            "Loi thay doi trang thai tai khoan nhan vien."
        )
        flash("Khong the cap nhat trang thai tai khoan.", "error")
    else:
        message = (
            "Da mo khoa tai khoan nhan vien."
            if employee.is_active
            else "Da khoa tai khoan nhan vien."
        )
        flash(message, "success")

    return redirect(url_for("admin.employee_list"))


@admin_bp.route("/employees/<int:employee_id>/edit", methods=["GET", "POST"])
@login_required
def employee_edit(employee_id):
    """Cho phép admin sửa thông tin nhân viên."""
    admin_required()

    employee = db.session.get(User, employee_id)

    if employee is None or employee.role != "employee":
        abort(404)

    if request.method == "POST":
        full_name = request.form.get("full_name", "").strip()
        email = request.form.get("email", "").strip().lower()
        phone = request.form.get("phone", "").strip()

        if not full_name or len(full_name) > 150:
            flash("Vui lòng nhập họ tên hợp lệ.", "error")
            return render_template(
                "admin/employee_edit.html",
                employee=employee,
            )

        if (
            not email
            or len(email) > 150
            or "@" not in email
            or any(char.isspace() for char in email)
        ):
            flash("Email không hợp lệ.", "error")
            return render_template(
                "admin/employee_edit.html",
                employee=employee,
            )

        if phone and len(phone) > 20:
            flash("Số điện thoại không được quá 20 ký tự.", "error")
            return render_template(
                "admin/employee_edit.html",
                employee=employee,
            )

        existing_user = db.session.scalar(
            db.select(User).where(
                db.func.lower(User.email) == email,
                User.id != employee.id,
            )
        )

        if existing_user:
            flash("Email đã được sử dụng bởi tài khoản khác.", "error")
            return render_template(
                "admin/employee_edit.html",
                employee=employee,
            )

        employee.full_name = full_name
        employee.email = email
        employee.phone = phone or None

        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
            admin_bp.logger.exception(
                "Lỗi cập nhật thông tin nhân viên."
            )
            flash("Không thể cập nhật thông tin nhân viên.", "error")
            return render_template(
                "admin/employee_edit.html",
                employee=employee,
            )

        log_action(
            action="UPDATE",
            target_type="employee",
            target_id=employee.id,
            description=f"Cập nhật thông tin nhân viên {employee.email}",
        )

        db.session.commit()

        flash("Cập nhật thông tin nhân viên thành công.", "success")
        return redirect(url_for("admin.employee_list"))

    return render_template(
        "admin/employee_edit.html",
        employee=employee,
    )


@admin_bp.route("/customers")
@login_required
def customer_list():
    """Hiển thị và tìm kiếm danh sách khách hàng."""
    admin_required()

    keyword = request.args.get("keyword", "").strip()

    query = User.query.filter_by(role="customer")

    if keyword:
        search_keyword = f"%{keyword}%"
        query = query.filter(
            db.or_(
                User.full_name.ilike(search_keyword),
                User.email.ilike(search_keyword),
                User.phone.ilike(search_keyword),
            )
        )

    customers = query.order_by(User.id.desc()).all()

    return render_template(
        "admin/customers.html",
        customers=customers,
        keyword=keyword,
    )


@admin_bp.route("/customers/<int:customer_id>")
@login_required
def customer_detail(customer_id):
    """Hiển thị thông tin và lịch sử đơn hàng khách hàng."""
    admin_required()

    customer = db.session.get(User, customer_id)

    if customer is None or customer.role != "customer":
        abort(404)

    orders = []

    if customer.email:
        orders = (
            Order.query
            .filter_by(user_id=customer.id)
            .order_by(Order.created_at.desc())
            .all()
        )

    return render_template(
        "admin/customer_detail.html",
        customer=customer,
        orders=orders,
    )


@admin_bp.route(
    "/customers/<int:customer_id>/toggle-active",
    methods=["POST"],
)
@login_required
def customer_toggle_active(customer_id):
    """Khóa hoặc mở khóa tài khoản khách hàng."""
    admin_required()

    customer = db.session.get(User, customer_id)

    if customer is None or customer.role != "customer":
        abort(404)

    customer.is_active = not customer.is_active

    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        flash("Không thể cập nhật trạng thái khách hàng.", "error")
        return redirect(url_for("admin.customer_list"))

    status_text = (
        "mở khóa" if customer.is_active else "khóa"
    )

    log_action(
        action="UPDATE",
        target_type="customer",
        target_id=customer.id,
        description=(
            f"Admin đã {status_text} tài khoản khách hàng "
            f"{customer.email}"
        ),
    )
    db.session.commit()

    flash(
        "Đã cập nhật trạng thái khách hàng thành công.",
        "success",
    )

    return redirect(url_for("admin.customer_list"))

@admin_bp.route("/reports/revenue/export")
@login_required
def export_revenue_report():
    """Xuất báo cáo doanh thu và chi tiết thanh toán ra Excel."""
    admin_required()
    from datetime import date, datetime, time, timedelta

    period = request.args.get("period", "").strip()
    start_text = request.args.get("start_date", "").strip()
    end_text = request.args.get("end_date", "").strip()

    today = date.today()
    start_date = None
    end_date = None

    # Xử lý bộ lọc nhanh.
    if period == "today":
        start_date = today
        end_date = today
    elif period == "last7":
        start_date = today - timedelta(days=6)
        end_date = today
    elif period == "this_month":
        start_date = today.replace(day=1)
        end_date = today
    elif period == "this_year":
        start_date = today.replace(month=1, day=1)
        end_date = today
    elif period:
        return "Bộ lọc thời gian không hợp lệ.", 400
    else:
        # Xử lý khoảng ngày người dùng nhập.
        try:
            if start_text:
                start_date = datetime.strptime(
                    start_text, "%Y-%m-%d"
                ).date()

            if end_text:
                end_date = datetime.strptime(
                    end_text, "%Y-%m-%d"
                ).date()
        except ValueError:
            return "Ngày lọc không hợp lệ.", 400

        if start_date and end_date and start_date > end_date:
            return "Ngày bắt đầu không được sau ngày kết thúc.", 400

    # Lấy các khoản thanh toán thành công, loại đơn đã hủy.
    query = (
        Payment.query
        .join(Order, Payment.order_id == Order.id)
        .filter(
            Payment.payment_status == "paid",
            Payment.paid_at.isnot(None),
            Order.order_status != "Đã hủy",
        )
    )

    # Lọc theo ngày thanh toán.
    if start_date:
        start_datetime = datetime.combine(start_date, time.min)
        query = query.filter(Payment.paid_at >= start_datetime)

    if end_date:
        end_datetime = datetime.combine(
            end_date + timedelta(days=1),
            time.min,
        )
        query = query.filter(Payment.paid_at < end_datetime)

    payments = query.order_by(Payment.paid_at.asc()).all()

    # Tạo file Excel.
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bao cao doanh thu"

    sheet.merge_cells("A1:F1")
    sheet["A1"] = "BÁO CÁO DOANH THU MENSTYLE"
    sheet["A1"].font = Font(
        name="Arial",
        size=16,
        bold=True,
        color="FFFFFF",
    )
    sheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="198754",
    )
    sheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    sheet.row_dimensions[1].height = 30

    # Thông tin bộ lọc.
    if start_date and end_date:
        filter_text = (
            f"Từ ngày {start_date.strftime('%d/%m/%Y')} "
            f"đến ngày {end_date.strftime('%d/%m/%Y')}"
        )
    elif start_date:
        filter_text = f"Từ ngày {start_date.strftime('%d/%m/%Y')}"
    elif end_date:
        filter_text = f"Đến ngày {end_date.strftime('%d/%m/%Y')}"
    else:
        filter_text = "Tất cả thời gian"

    sheet.merge_cells("A2:F2")
    sheet["A2"] = f"Khoảng thời gian: {filter_text}"

    total_revenue = sum(
        (payment.amount or 0 for payment in payments),
        start=0,
    )
    paid_order_count = len({
        payment.order_id for payment in payments
    })

    sheet["A4"] = "Tổng doanh thu"
    sheet["B4"] = float(total_revenue)
    sheet["A5"] = "Số đơn có thanh toán"
    sheet["B5"] = paid_order_count

    sheet["B4"].number_format = '#,##0 "đ"'

    # Tiêu đề bảng chi tiết.
    headers = [
        "Mã đơn hàng",
        "Phương thức thanh toán",
        "Số tiền (VNĐ)",
        "Thời điểm thanh toán",
        "Trạng thái đơn",
        "Mã thanh toán",
    ]

    header_row = 7

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="292E56",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # Ghi dữ liệu.
    for row_number, payment in enumerate(payments, start=header_row + 1):
        order = payment.order

        sheet.cell(
            row=row_number,
            column=1,
            value=order.order_code if order else "",
        )
        sheet.cell(
            row=row_number,
            column=2,
            value=payment.payment_method or "",
        )

        amount_cell = sheet.cell(
            row=row_number,
            column=3,
            value=float(payment.amount or 0),
        )
        amount_cell.number_format = '#,##0 "đ"'

        paid_at_cell = sheet.cell(
            row=row_number,
            column=4,
            value=payment.paid_at,
        )
        paid_at_cell.number_format = "dd/mm/yyyy hh:mm"

        sheet.cell(
            row=row_number,
            column=5,
            value=order.order_status if order else "",
        )
        sheet.cell(
            row=row_number,
            column=6,
            value=payment.transaction_code or "",
        )

    # Định dạng bảng.
    sheet.freeze_panes = "A8"
    sheet.auto_filter.ref = (
        f"A7:F{max(header_row + len(payments), header_row)}"
    )

    column_widths = {
        "A": 22,
        "B": 25,
        "C": 20,
        "D": 24,
        "E": 22,
        "F": 25,
    }

    for column, width in column_widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows(
        min_row=8,
        max_row=sheet.max_row,
        min_col=1,
        max_col=6,
    ):
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    # Xuất file để người dùng tải xuống.
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    workbook.close()

    return send_file(
        output,
        as_attachment=True,
        download_name="bao_cao_doanh_thu_menstyle.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
